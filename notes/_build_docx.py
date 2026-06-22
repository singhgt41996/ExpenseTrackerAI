"""Generate a nicely formatted Word document of the Safe Area & Status Bar notes."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x2E, 0x86, 0xDE)      # blue
CODE_BG = "F2F3F5"
CODE_TEXT = RGBColor(0x24, 0x29, 0x2E)
GREY = RGBColor(0x55, 0x5B, 0x66)

doc = Document()

# Base font
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        shade(p, CODE_BG)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = CODE_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    _add_rich(p, text)
    return p


def _add_rich(p, text):
    """Render **bold** and `code` inline."""
    import re
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`") and t.endswith("`"):
            r = p.add_run(t[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xC0, 0x34, 0x4D)
        else:
            p.add_run(t)


def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text); r.font.color.rgb = ACCENT
    return p


def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    return p


def para(text):
    p = doc.add_paragraph()
    _add_rich(p, text)
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].add_run(htext).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            _add_rich(cells[i].paragraphs[0], val)
    doc.add_paragraph()


# ---------------- Title ----------------
title = doc.add_heading(level=0)
trun = title.add_run("React Native Notes — Safe Area & Status Bar")
sub = doc.add_paragraph()
srun = sub.add_run("Personal revision notes from the FinFlow project")
srun.italic = True
srun.font.color.rgb = GREY
doc.add_paragraph()

# ---------------- 1 ----------------
h1("1. The JSX / React import rule")
add_bullet("With `\"jsx\": \"react-native\"` in `tsconfig.json` (the **classic** JSX transform), every JSX tag compiles to `React.createElement(...)`, so **`React` must be imported** in the file.")
add_bullet("Error \"'React' refers to a UMD global, but the current file is a module\" = you used JSX but forgot to import React.")
add_bullet("Fix:")
add_code("import React, { useEffect, useState } from 'react';")
add_bullet("Modern alternative: set `\"jsx\": \"react-jsx\"` (automatic runtime) → no React import needed anywhere. But unused `import React` lines elsewhere may then be flagged by `noUnusedLocals`.")

# ---------------- 2 ----------------
h1("2. SafeAreaProvider")
add_bullet("It is **only a context provider**. It measures the device's safe-area insets (notch, status bar, home indicator) and makes those numbers available to children.")
add_bullet("**It does NOT add any padding by itself.** Wrapping your app in it changes nothing visually.")
add_bullet("Put it **once at the root** of the app.")
add_bullet("**Gotcha:** the component that renders `SafeAreaProvider` cannot use `useSafeAreaInsets()` itself — the hook / `SafeAreaView` must live in a child **nested under** the provider. Otherwise insets return 0.")
add_code("""function App() {
  return (
    <SafeAreaProvider>
      <AppContent />   {/* hook works here, not in App */}
    </SafeAreaProvider>
  );
}""")

# ---------------- 3 ----------------
h1("3. Two ways to APPLY the insets")
h2("a) SafeAreaView (declarative — preferred for containers)")
add_bullet("A `View` that auto-applies insets as padding.")
add_bullet("Use the one from **react-native-safe-area-context**, NOT the deprecated one from react-native.")
add_bullet("`edges` prop chooses which sides to pad: ['top'], ['bottom'], ['top','bottom'], etc.")
add_code("""<SafeAreaView style={{ flex: 1 }} edges={['top', 'bottom']}>
  {children}
</SafeAreaView>""")
h2("b) useSafeAreaInsets() (hook — for raw numbers / special cases)")
add_bullet("Returns { top, bottom, left, right } as numbers.")
add_bullet("Use when you need the value for something other than simple container padding:")
add_bullet("Floating button → `marginBottom: insets.bottom`", level=1)
add_bullet("ScrollView → `contentContainerStyle={{ paddingBottom: insets.bottom }}`", level=1)
add_bullet("Custom header, animations, layout math", level=1)
para("**Rule of thumb:** `SafeAreaView` for the screen container; the hook for special elements inside.")

# ---------------- 4 ----------------
h1("4. Screen Wrapper Pattern (reusable)")
add_bullet("Don't repeat safe-area logic on every screen. Make one reusable wrapper (a \"template\" in atomic design).")
add_bullet("`SafeAreaProvider` stays at root; each screen wraps its content in `<Screen>`.")
add_bullet("The wrapper does NOT go in the navigator config — it goes **inside each screen component** as the root element.")
add_code("""// src/components/templates/Screen/index.tsx
import { SafeAreaView, Edge } from 'react-native-safe-area-context';

type ScreenProps = {
  children: React.ReactNode;
  edges?: Edge[];
  backgroundColor?: string;
  barStyle?: 'dark-content' | 'light-content';
};

export const Screen = ({
  children,
  edges = ['top', 'bottom'],
  backgroundColor = '#fff',
  barStyle = 'dark-content',
}: ScreenProps) => (
  <SafeAreaView style={{ flex: 1, backgroundColor }} edges={edges}>
    <StatusBar barStyle={barStyle} backgroundColor={backgroundColor} />
    {children}
  </SafeAreaView>
);""")
para("How it's used with navigation:")
add_code("""// Navigator = routing only (no wrappers here)
<Stack.Navigator screenOptions={{ headerShown: false }}>
  <Stack.Screen name="Login" component={LoginScreen} />
</Stack.Navigator>

// Screen = uses the wrapper internally
const LoginScreen = () => (
  <Screen edges={['top', 'bottom']}>
    {/* UI */}
  </Screen>
);""")

# ---------------- 5 ----------------
h1("5. Safe Area + React Navigation (important)")
add_bullet("Navigation **headers and bottom tab bars handle safe area automatically.**")
add_bullet("If a screen **has a header** → wrapper should use `edges={['bottom']}` (header covers top; padding top too = double gap).")
add_bullet("If a screen **has no header** (e.g. auth screens with headerShown: false) → use `edges={['top', 'bottom']}`.")
add_bullet("This is why `edges` should be a prop.")

# ---------------- 6 ----------------
h1("6. StatusBar  <StatusBar />")
add_bullet("Controls the OS top bar (clock, battery, wifi).")
add_bullet("It's **imperative** — doesn't render UI in place; it tells the OS how the bar looks.")
add_bullet("React Native merges all mounted `StatusBar` components; **the last one to mount/update wins.**")
h2("Props")
table(
    ["Prop", "Meaning"],
    [
        ["barStyle", "Color of icons/text: `dark-content` = dark icons (light bg), `light-content` = white icons (dark bg)"],
        ["backgroundColor", "Background behind the bar — **Android only**"],
        ["translucent", "Android: draw under the status bar"],
        ["hidden", "Hide the bar"],
    ],
)
para("Naming trick: `dark-content` = dark text, used on a light background.")
h2("Root vs per-page")
add_bullet("Put one at the **root** as the default.")
add_bullet("Add a `<StatusBar>` **per-screen only when that screen needs something different.**")
h2("Making status bar color match the page")
add_bullet("**iOS:** status bar is always a transparent overlay — no `backgroundColor`. Color the **top safe-area region** instead (e.g. SafeAreaView backgroundColor), and set `barStyle` for readable icons.")
add_bullet("**Android:** set `backgroundColor` on `<StatusBar>` directly.")
add_bullet("**Cross-platform:** set `barStyle` + `backgroundColor` (Android) + color the top safe area (iOS). Bundling this into the Screen wrapper handles both at once.")
h2("Caveat with navigation")
add_bullet("Because last-mounted wins, navigating **back** may not re-apply a screen's status bar. Fix by re-applying on focus (`useFocusEffect` + `StatusBar.setBarStyle(...)`) or keeping `<StatusBar>` inside each screen.")

# ---------------- Summary ----------------
h1("Quick Summary Table")
table(
    ["Thing", "Job"],
    [
        ["SafeAreaProvider", "Measures insets; provides numbers. No padding itself. Root, once."],
        ["SafeAreaView", "Applies insets as padding on chosen `edges`. Use for containers."],
        ["useSafeAreaInsets()", "Raw inset numbers for special elements."],
        ["Screen wrapper", "Reusable: SafeAreaView + StatusBar + bg color. Used inside each screen."],
        ["StatusBar", "Tells OS how the top bar looks. Root default + per-page overrides."],
    ],
)

out = "notes/SafeArea-and-StatusBar-Notes.docx"
doc.save(out)
print("Saved:", out)
